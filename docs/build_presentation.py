"""Generate editable PowerPoint slides and a matching PDF handout."""
import ast
import json
import sqlite3
import subprocess
import sys
import tempfile
import textwrap
from pathlib import Path
from PIL import Image

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parents[1]
W, H = 960, 540
BG, PANEL, WHITE, MUTED, TEAL, ORANGE = '101B2B', '1B2B40', 'F4F7FB', 'B7C6D8', '52D5C4', 'FFBC75'
FONTDIR = Path('/System/Library/Fonts/Supplemental')
pdfmetrics.registerFont(TTFont('Arial', str(FONTDIR / 'Arial.ttf')))
pdfmetrics.registerFont(TTFont('ArialBold', str(FONTDIR / 'Arial Bold.ttf')))
pdfmetrics.registerFont(TTFont('Code', str(FONTDIR / 'Courier New.ttf')))
prs = Presentation()
prs.slide_width, prs.slide_height = Pt(W), Pt(H)
pdf = canvas.Canvas(str(ROOT / 'ProjectPresentation.pdf'), pagesize=(W, H))
pdf.setTitle('Airline Reservation System — Project Presentation')
notes = []
slide = None

def rect(x, y, w, h, color=PANEL):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    pdf.setFillColorRGB(*[int(color[i:i+2], 16) / 255 for i in (0, 2, 4)])
    pdf.rect(x, H-y-h, w, h, stroke=0, fill=1)

def text(x, y, w, content, size=20, color=WHITE, bold=False, code=False):
    lines = content.split('\n')
    height = (size * 1.28) * len(lines) + 10
    assert x >= 0 and x + w <= W and y + height < H, (content, y, height)
    box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = line
        para.font.name = 'Courier New' if code else 'Arial'
        para.font.size = Pt(size)
        para.font.bold = bold
        para.font.color.rgb = RGBColor.from_string(color)
        para.space_before = para.space_after = Pt(0)
        para.line_spacing = Pt(size * 1.28)
    pdf.setFillColorRGB(*[int(color[i:i+2],16)/255 for i in (0,2,4)])
    font = 'Code' if code else ('ArialBold' if bold else 'Arial')
    pdf.setFont(font, size)
    for i,line in enumerate(lines):
        assert pdfmetrics.stringWidth(line, font, size) <= w + 2, ('Text too wide', line)
        pdf.drawString(x, H-y-size-i*size*1.28, line)

def new(title, subtitle, presenter, speech):
    global slide
    if slide is not None:
        pdf.showPage()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    pdf.setFillColorRGB(16/255,27/255,43/255)
    pdf.rect(0,0,W,H,fill=1,stroke=0)
    n=len(prs.slides)
    text(42,22,870,'AIRLINE RESERVATION SYSTEM',11,TEAL,bold=True)
    text(42,57,880,title,32,bold=True)
    text(42,106,880,subtitle,16,MUTED)
    rect(42,493,876,1,'35455A')
    text(889,506,40,f'{n:02d}',11,TEAL)
    slide.notes_slide.notes_text_frame.text = speech
    notes.append(f'## Slide {n}: {title}\n\nPresenter: {presenter}\n\n{speech}\n')

def codebox(x,y,w,h,title,content,size=15):
    rect(x,y,w,h)
    text(x+16,y+12,w-32,title,12,TEAL,bold=True)
    text(x+16,y+42,w-32,content,size,WHITE,code=True)

def excerpt(file, method):
    source=(ROOT/file).read_text()
    tree=ast.parse(source)
    node=next(n for n in ast.walk(tree) if isinstance(n,ast.FunctionDef) and n.name==method)
    return textwrap.dedent('\n'.join(source.splitlines()[node.lineno-1:node.end_lineno]))

def screenshot(time, x=42, y=155, w=490, h=310, bottom=False):
    path=ROOT/'docs/screenshots'/f'Screenshot 2026-10-18 at {time}.png'
    im=Image.open(path)
    if bottom:
        im=im.crop((0, im.height-285, im.width, im.height))
        path=ROOT/'docs/report_assets/booking_detail.png'
        im.save(path)
    iw,ih=im.size
    scale=min(w/iw,h/ih)
    dw,dh=iw*scale,ih*scale
    slide.shapes.add_picture(str(path),Pt(x),Pt(y),width=Pt(dw),height=Pt(dh))
    pdf.drawImage(str(path),x,H-y-dh,width=dw,height=dh)

def demo(title,time,method,file,caption,speech,bottom=False):
    presenter='Kushekbayev Batyrlan' if file=='main.py' else 'Abduldinov Alibek'
    new(title,'Actual terminal screenshot • code excerpt • observed result',presenter,speech)
    screenshot(time,bottom=bottom)
    snippet=excerpt(file,method)
    lines=snippet.splitlines()
    if len(lines)>12:
        lines=lines[:11]+['    # ... remaining implementation']
    # Wrap long source lines for display without changing the source file.
    display=[]
    for line in lines:
        display.extend(textwrap.wrap(line, width=45, subsequent_indent='        ',
            replace_whitespace=False,drop_whitespace=False) or [''])
    if len(display)>14:
        display=display[:13]+['    # ... excerpt continues']
    codebox(555,155,363,260,file+' / '+method+'()', '\n'.join(display),11)
    text(555,433,363,caption,14,TEAL)

exec(compile((ROOT / 'docs/presentation_content.py').read_text(), str(ROOT / 'docs/presentation_content.py'), 'exec'))
